"""v3 of the SETU submission deck.

Revisions over v2:
  · slide 1 — Problem Statement ID and Theme filled in
  · slide 2 — rewritten in plain language: the slide now answers "what is your
    solution?" without acronyms, and the chart is flattened and explained rather
    than left to speak for itself
  · slide 3 — client-runtime lane untouched; the sync and cloud planes get tighter
    node padding, their own colour, and larger, coloured arrows
  · slide 4 — risk table type enlarged to use the space it occupies

Reuses every primitive and slides 5-6 from build_deck_v2.

    python assets/make_figures_v2.py && python assets/build_deck_v3.py
"""
import os
import build_deck_v2 as v2
from build_deck_v2 import (rect, line, arrow, tb, picture, heading, kill,
                           set_title, set_oval, MEMBERS, TEAM,
                           NAVY, BLUE, SAFFRON, GREEN, RED, INK, MUTED, FIELD,
                           RULE, PANEL, PANEL2, LANE1, WHITE, PALE,
                           SANS, SERIF, MONO, TOP, BOT, L, R, W)
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "SIH2026-IDEA-Presentation-Format.pptx")
DST = os.path.join(ROOT, "SETU-SIH2026-Idea-Presentation-v3.pptx")

PS_ID = "SIH26168"
THEME = "Smart Vehicles"

SYNCBG = RGBColor(0xE7, 0xF0, 0xF9)     # network plane
CLOUDBG = RGBColor(0xFB, 0xF2, 0xE6)    # compute plane


# =============================================================== SLIDE 1
def slide1(sl):
    kill(sl, "TextBox 9")
    for sh in sl.shapes:
        if sh.name.startswith("Subtitle") and sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = ""
            tf = sh.text_frame
            p = tf.paragraphs[1] if len(tf.paragraphs) > 1 else tf.paragraphs[0]
            r = p.add_run(); r.text = "SETU · INTELLIGENT DEAD RECKONING"
            r.font.size = Pt(23); r.font.bold = True
            r.font.color.rgb = NAVY; r.font.name = SERIF

    F = 19.0
    blank = ("____________", F, False, FIELD)
    rows = [
        [("Problem Statement ID –  ", F, True, NAVY), (PS_ID, F, False, INK)],
        [("Problem Statement Title-  ", F, True, NAVY),
         ("AI-ML based Intelligent Dead Reckoning system for seamless navigation",
          F, False, INK)],
        [("Theme-  ", F, True, NAVY), (THEME, F, False, INK)],
        [("PS Category-  ", F, True, NAVY), ("Software", F, False, INK)],
        [("Team ID-  ", F, True, NAVY), blank],
        [("Team Name (Registered on portal)-  ", F, True, NAVY), (TEAM, F, False, INK)],
    ]
    tb(sl, L, 2.10, 6.70, 4.60,
       [{"runs": r, "sa": 9, "line": 1.0} for r in rows], pad=0.0)

    cx, cy, cw, ch = 7.12, 5.60, 5.79, 1.26
    rect(sl, cx, cy, cw, ch, fill=WHITE, line=NAVY, lw=1.25)
    tb(sl, cx, cy + 0.06, cw - 1.72, ch - 0.12, [
        {"runs": [("Seamless Egomotion Tracking under Unavailable-GNSS", 10.5, True, NAVY)],
         "sa": 3, "line": 1.0},
        {"runs": [("Lane-level position through tunnels, underpasses, basement parking and "
                   "urban canyons — from a smartphone IMU alone.", 9.6, False, INK)],
         "line": 1.05}], pad=0.12)
    line(sl, cx + cw - 1.74, cy + 0.16, cx + cw - 1.74, cy + ch - 0.16, RULE, 1.0)
    tb(sl, cx + cw - 1.70, cy + 0.06, 1.62, ch - 0.12, [
        {"runs": [("≈1%", 28, True, GREEN)], "align": PP_ALIGN.CENTER, "sa": 0, "line": 0.92},
        {"runs": [("drift per km, against\na 10 % ceiling", 8.6, False, MUTED)],
         "align": PP_ALIGN.CENTER, "line": 1.04}], anchor="m", pad=0.03)

    rect(sl, 0, 6.95, 13.333, 0.55, fill=BLUE)
    runs = [("TEAM  ", 11, True, PALE), (TEAM, 12, True, WHITE), ("      ", 11, False, WHITE)]
    for i, (name, role) in enumerate(MEMBERS):
        if i:
            runs.append(("  ·  ", 10.5, False, PALE))
        runs.append((name, 10.5, True, WHITE))
        if role:
            runs.append((" (" + role + ")", 10.5, False, PALE))
    tb(sl, 0.40, 6.95, 12.53, 0.55,
       [{"runs": runs, "align": PP_ALIGN.CENTER, "line": 1.0}], anchor="m")


# =============================================================== SLIDE 2
def slide2(sl):
    kill(sl, "TextBox 8"); set_oval(sl)
    set_title(sl, "SETU — NAVIGATION WITHOUT GNSS", 32)

    tb(sl, L, 1.30, W, 0.64, [{"runs": [
        ("When GPS disappears, today's apps freeze, jump, or call the turn too late. ",
         18, False, INK),
        ("SETU keeps your position moving correctly on the map, using nothing but the sensors "
         "already inside the phone.", 18, True, NAVY)], "line": 1.06}], pad=0.0)
    line(sl, L, 1.98, R, 1.98, RULE, 1.0)

    LW, RX, RW = 5.10, 5.76, 7.15

    # ---------------- left: what it is, in plain words
    heading(sl, L, 2.04, LW, "Detailed explanation of the proposed solution", 12.5)
    rect(sl, L, 2.48, LW, 0.46, fill=PANEL, line=RULE)
    tb(sl, L, 2.48, LW, 0.46, [{"runs": [
        ("An Android app and a small on-device engine.", 10.4, True, NAVY),
        ("  No dongle, no extra hardware, no internet.", 10.0, False, INK)],
        "line": 1.0}], anchor="m", pad=0.11)

    ideas = [
        (GREEN, "The phone hears the wheels.",
         "Spinning wheels shake the car, and the shake speeds up as you do. The app reads that "
         "pitch, so it measures speed instead of guessing it from acceleration — the way a "
         "mechanic listens to an engine."),
        (SAFFRON, "Every turn re-checks the speed.",
         "In a bend, the sideways push divided by how sharply you are turning gives the exact "
         "speed — free calibration several times a minute, and it works on a leaning "
         "two-wheeler too."),
        (BLUE, "The road becomes the ruler.",
         "The phone knows how much it has turned; the offline map knows the shape of every "
         "road. Line the two up and your position falls into place with no satellites at all."),
    ]
    for i, (col, label, body) in enumerate(ideas):
        yy = 3.02 + i * 1.02
        rect(sl, L, yy, LW, 0.96, fill=WHITE, line=col, lw=1.1)
        tb(sl, L, yy + 0.04, LW, 0.88, [
            {"runs": [(label, 11.5, True, col)], "sa": 2, "line": 1.0},
            {"runs": [(body, 10.0, False, INK)], "line": 1.04}], pad=0.11)

    # ---------------- right: the chart, flattened, then read out loud
    heading(sl, RX, 2.04, RW, "How it addresses the problem", 12.5)
    b = picture(sl, RX, 2.48, RW, os.path.join(HERE, "fig-error-law-flat.png"))

    tb(sl, RX, b + 0.10, RW, 0.24, [{"runs": [
        ("The graph answers one question: after a kilometre with no GPS, how far off are you?",
         10.4, True, NAVY)], "line": 1.0}], pad=0.0)
    readout = [
        (RED, "Pure inertial · 180 m",
         "acceleration simply added up — tiny errors compound, so you are a block away "
         "after one minute.", 0.38),
        (SAFFRON, "Best published phone result · 31 m",
         "error grows steadily with distance, 3.1 % of every kilometre.", 0.21),
        (GREEN, "SETU · 8.5 m",
         "the line drops each time we recognise something — a bend whose shape matches the "
         "map, or a tunnel's magnetic signature. Error is reset instead of piling up.", 0.38),
    ]
    ry = b + 0.40
    for col, label, body, h in readout:
        rect(sl, RX + 0.01, ry + 0.055, 0.17, 0.10, fill=col)
        tb(sl, RX + 0.24, ry, RW - 0.24, h + 0.04, [{"runs": [
            (label + "  —  ", 10.2, True, col), (body, 10.2, False, INK)], "line": 1.02}],
           pad=0.0)
        ry += h + 0.05

    # ---------------- innovation, one line each
    heading(sl, L, 6.06, W, "Innovation and uniqueness of the solution", 12.5)
    nov = [("N1", GREEN, "Speed read from frequency"),
           ("N2", SAFFRON, "Two-wheeler lean physics"),
           ("N3", BLUE, "Scale-free map registration"),
           ("N4", NAVY, "CAN wheel speeds as teacher"),
           ("N5", RED, "Learned σ and Q in one filter")]
    cw = (W - 4 * 0.12) / 5
    for i, (tag, col, head) in enumerate(nov):
        gx = L + i * (cw + 0.12)
        rect(sl, gx, 6.44, cw, 0.38, fill=PANEL, line=RULE)
        tb(sl, gx, 6.44, cw, 0.38, [{"runs": [
            (tag + "  ", 9.8, True, col), (head, 9.8, True, INK)],
            "align": PP_ALIGN.CENTER, "line": 1.0}], anchor="m", pad=0.05)


# =============================================================== SLIDE 3
def slide3(sl):
    kill(sl, "TextBox 8"); set_oval(sl)
    set_title(sl, "TECHNICAL APPROACH")
    heading(sl, L, TOP, W, "Methodology and process for implementation")

    GX, GW = L, 1.16
    NX = GX + GW + 0.10
    NW = R - NX
    HA, HAW, HB, HBW = NX, 5.30, 7.30, 5.61

    # ---- lane frames
    for ly, lh, n1, n2, s1, s2, scol, bcol, fillc, blw in [
        (1.70, 1.86, "CLIENT", "RUNTIME", "offline", "zero network calls", GREEN, NAVY, LANE1, 1.5),
        (3.82, 0.74, "SYNC", "PLANE", "out-of-band", "async, never blocking", BLUE, BLUE, SYNCBG, 1.25),
        (4.82, 1.00, "CLOUD", "PLANE", "batch", "region-sharded", SAFFRON, SAFFRON, CLOUDBG, 1.25),
    ]:
        rect(sl, GX, ly, W, lh, fill=fillc, line=bcol, lw=blw)
        tb(sl, GX, ly, GW, lh, [
            {"runs": [(n1, 10.5, True, NAVY)], "sa": 0, "line": 0.92},
            {"runs": [(n2, 10.5, True, NAVY)], "sa": 3, "line": 0.92},
            {"runs": [(s1, 8.0, True, scol)], "sa": 0, "line": 0.95},
            {"runs": [(s2, 8.0, False, MUTED)], "line": 0.95},
        ], anchor="m", pad=0.09)
        line(sl, GX + GW, ly + 0.08, GX + GW, ly + lh - 0.08, RULE, 0.9)

    # ---- lane 1: unchanged from v2
    groups = [
        (NX, 1.70, "SENSE", [("IMU", "200–400 Hz"), ("GNSS raw", "NavIC L5"),
                             ("Baro", "Magnetometer"), ("External IMU", "MEMS / FOG")]),
        (3.64, 1.70, "CONDITION", [("STFT", "harmonic ridge"), ("Notch", "engine order"),
                                   ("Shock", "pothole reject"), ("ACE", "mount solve")]),
        (5.60, 3.29, "OBSERVE", [("SVO", "spectral speed"), ("CTS", "turn speed"),
                                 ("CSA", "map arc length"), ("EFA", "field anchors"),
                                 ("NHC", "zero side-slip"), ("GQM", "sat trust")]),
        (9.15, 1.85, "FUSE", [("RI-EKF", "SE₂(3)"), ("learned σ", "per measurement"),
                              ("learned Q", "transformer"), ("RB-PF", "road graph")]),
        (11.26, 1.57, "DELIVER", [("10 Hz", "phone pose"), ("200 Hz", "edge pose"),
                                  ("Lane-level", "+ tube"), ("API", "gRPC · NMEA")]),
    ]
    hy, cy0 = 1.76, 2.04
    for gx, gw, name, chips in groups:
        rect(sl, gx, hy, gw, 0.24, fill=NAVY)
        tb(sl, gx, hy, gw, 0.24, [{"runs": [(name, 8.6, True, WHITE)],
                                   "align": PP_ALIGN.CENTER}], anchor="m", pad=0.02)
        if name == "OBSERVE":
            sub = (gw - 0.10) / 2
            for j, (a, bb) in enumerate(chips):
                px = gx + (j % 2) * (sub + 0.10)
                py = cy0 + (j // 2) * 0.42
                rect(sl, px, py, sub, 0.36, fill=WHITE, line=RULE)
                tb(sl, px, py, sub, 0.36, [
                    {"runs": [(a, 9.0, True, NAVY)], "sa": 0, "line": 0.94},
                    {"runs": [(bb, 7.8, False, MUTED)], "line": 0.94}], anchor="m", pad=0.05)
        else:
            for j, (a, bb) in enumerate(chips):
                py = cy0 + j * 0.34
                rect(sl, gx, py, gw, 0.29, fill=WHITE, line=RULE)
                tb(sl, gx, py, gw, 0.29, [{"runs": [
                    (a, 8.8, True, NAVY), ("   " + bb, 7.8, False, MUTED)],
                    "line": 0.95}], anchor="m", pad=0.06)
    for ax in (3.40, 5.36, 8.91, 11.02):
        arrow(sl, ax, 2.67, 0.22, 0.20, BLUE)

    # ---- lane 2: sync plane, tighter nodes and its own colour
    for hx, hw, hname, hnote, hcol, nodes in [
        (HA, HAW, "DISTRIBUTION", "signed, immutable, versioned bundles — installed offline",
         BLUE, ["Model registry", "Map / tile CDN", "Config & flags"]),
        (HB, HBW, "CONTRIBUTION", "opt-in · ~1 KB per edge · signatures only, never trajectories",
         GREEN, ["API gateway", "Stream queue", "k-anon aggregator"]),
    ]:
        tb(sl, hx, 3.87, hw, 0.20, [{"runs": [
            (hname, 9.0, True, hcol), ("     " + hnote, 8.2, False, MUTED)],
            "line": 0.95}], pad=0.0)
        nw = (hw - 2 * 0.24) / 3
        for j, nm in enumerate(nodes):
            nx = hx + j * (nw + 0.24)
            rect(sl, nx, 4.10, nw, 0.36, fill=WHITE, line=hcol, lw=1.25)
            tb(sl, nx, 4.10, nw, 0.36, [{"runs": [(nm, 9.4, True, NAVY)],
                                         "align": PP_ALIGN.CENTER, "line": 0.96}],
               anchor="m", pad=0.03)

    # ---- lane 3: cloud plane
    tb(sl, HA, 4.87, HAW, 0.20, [{"runs": [
        ("MAP BUILD", 9.0, True, SAFFRON),
        ("     per region, independently", 8.2, False, MUTED)], "line": 0.95}], pad=0.0)
    nw = (HAW - 2 * 0.24) / 3
    for j, nm in enumerate(["OSM extract", "Graph + κ(s) LUT", "DEM · PMTiles"]):
        nx = HA + j * (nw + 0.24)
        rect(sl, nx, 5.10, nw, 0.36, fill=WHITE, line=SAFFRON, lw=1.25)
        tb(sl, nx, 5.10, nw, 0.36, [{"runs": [(nm, 9.4, True, NAVY)],
                                     "align": PP_ALIGN.CENTER, "line": 0.96}],
           anchor="m", pad=0.03)
        if j < 2:
            arrow(sl, nx + nw + 0.02, 5.15, 0.20, 0.26, SAFFRON)

    tb(sl, HB, 4.87, HBW, 0.20, [{"runs": [
        ("LEARNING", 9.0, True, SAFFRON),
        ("     nightly retrain, gated on the benchmark", 8.2, False, MUTED)],
        "line": 0.95}], pad=0.0)
    nw2 = (HBW - 3 * 0.24) / 4
    for j, nm in enumerate(["Data lake", "Training cluster", "Eval gate", "Telemetry & KPI"]):
        nx = HB + j * (nw2 + 0.24)
        rect(sl, nx, 5.10, nw2, 0.36, fill=WHITE, line=SAFFRON, lw=1.25)
        tb(sl, nx, 5.10, nw2, 0.36, [{"runs": [(nm, 9.4, True, NAVY)],
                                      "align": PP_ALIGN.CENTER, "line": 0.96}],
           anchor="m", pad=0.03)
        if j < 3:
            arrow(sl, nx + nw2 + 0.02, 5.15, 0.20, 0.26, SAFFRON)

    tb(sl, NX, 5.49, NW, 0.28, [{"runs": [
        ("SCALES BY SHARDING, NOT BY CAPACITY", 8.6, True, NAVY),
        ("     zero runtime API calls, so device count adds no serving load  ·  "
         "≤250 MB per region  ·  immutable CDN bundles  ·  contribution ~1 KB per edge",
         8.4, False, INK)], "line": 1.0}], anchor="m", pad=0.0)

    # ---- inter-plane arrows: large and colour-coded by direction
    for ay in (3.56, 4.56):
        arrow(sl, 3.24, ay, 0.36, 0.26, BLUE, MSO_SHAPE.UP_ARROW)
        arrow(sl, 10.10, ay, 0.36, 0.26, GREEN, MSO_SHAPE.DOWN_ARROW)

    # ---- technologies
    heading(sl, L, 5.90, W, "Technologies to be used")
    tw = (W - 4 * 0.10) / 5
    for i, (name, body) in enumerate([
        ("Training", "PyTorch · Lightning · Hydra · MLflow · masked-IMU pretraining"),
        ("On-device", "LiteRT + XNNPACK · ExecuTorch · int8 QAT · 3.6 MB total"),
        ("Engine", "C++20 · Eigen · SE₂(3) Lie · FlatBuffers · GTest replay"),
        ("Mobile", "Kotlin · Compose · MapLibre + PMTiles · SensorDirectChannel"),
        ("Edge & maps", "aarch64 · gRPC / NMEA / ROS 2 · osmium · CartoDEM"),
    ]):
        gx = L + i * (tw + 0.10)
        rect(sl, gx, 6.28, tw, 0.54, fill=PANEL, line=RULE)
        tb(sl, gx, 6.30, tw, 0.50, [
            {"runs": [(name, 9.0, True, NAVY)], "sa": 1, "line": 0.96},
            {"runs": [(body, 8.2, False, INK)], "line": 1.0}], pad=0.08)


# =============================================================== SLIDE 4
def slide4(sl):
    kill(sl, "TextBox 8"); set_oval(sl)
    set_title(sl, "FEASIBILITY AND VIABILITY")

    LW = 4.34
    y = heading(sl, L, TOP, LW, "Analysis of the feasibility of the idea")
    for i, (claim, ev) in enumerate([
        ("Data already in hand", "IO-VNBD: 5,700 km, phone IMU and CAN bus recorded together"),
        ("Core physics is closed-form", "v = a_lat/Ω and f = v/2πR are algebra; only residuals learn"),
        ("Fits today's phones", "3.6 MB int8, 2.7 ms per tick, CPU only — no NPU needed"),
        ("Sensors already present", "IMU, magnetometer, barometer, raw GNSS with NavIC L5"),
        ("Maps fit offline", "≤250 MB per metro region, no connectivity at runtime"),
        ("Nothing to buy or fit", "no dongle, no encoder, no roadside infrastructure"),
    ]):
        yy = y + i * 0.44
        tb(sl, L, yy, LW, 0.42, [
            {"runs": [(claim, 10.4, True, NAVY)], "sa": 0, "line": 0.98},
            {"runs": [(ev, 9.2, False, INK)], "line": 1.0}], pad=0.0)

    b = picture(sl, L, 4.42, LW, os.path.join(HERE, "fig-tiers.png"))
    tb(sl, L, b + 0.04, LW, 0.30, [{"runs": [
        ("Every device class clears the ceiling. IO-VNBD's 10 Hz stream is Tier C.",
         9.0, False, MUTED)], "line": 1.0}], pad=0.0)

    tx, tw_ = 5.00, 7.91
    rows = [
        ("CSA ambiguous on featureless or badly-mapped roads", "High", RED,
         "Landmark saliency scoring · top-k hypotheses in the particle filter · honest σ. "
         "Falls back to classical HMM map matching."),
        ("Axle harmonics weak — smooth road, quiet EV, soft mount", "Medium", SAFFRON,
         "Mount-quality gate · summation over 8+ orders · learned σ down-weights it. "
         "Falls back to Tier B."),
        ("IO-VNBD phone stream is 10 Hz, so cannot validate SVO", "Known", BLUE,
         "Declared as capability tiers. SVO validated on self-collected 400 Hz logs."),
        ("No two-wheeler data in IO-VNBD — four cars, no Indian roads", "Known", BLUE,
         "Lean relation imposed as a physics loss, so few samples suffice. Dedicated collection."),
        ("Sensor rate throttled below 200 Hz on some devices", "Medium", SAFFRON,
         "Tier detected at runtime · SensorDirectChannel · verified on three device classes."),
        ("Uneven OpenStreetMap geometry on Indian roads", "Medium", SAFFRON,
         "Covariance scaled by OSM confidence · explicit off-road particle · provider adapter."),
    ]
    shp = sl.shapes.add_table(len(rows) + 1, 3, In(tx), In(TOP), In(tw_), In(4.84))
    t = shp.table
    t.first_row = False; t.horz_banding = False
    for c, wd in zip(t.columns, (2.76, 0.86, 4.29)):
        c.width = In(wd)
    t.rows[0].height = In(0.42)
    for i in range(1, len(rows) + 1):
        t.rows[i].height = In(0.736)
    for j, htxt in enumerate(["Potential challenges and risks", "Severity",
                              "Strategies for overcoming these challenges"]):
        cell = t.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        cell.margin_left = cell.margin_right = In(0.07)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = htxt
        r.font.size = Pt(10.6); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = SANS
        if j == 1:
            p.alignment = PP_ALIGN.CENTER
    for i, (risk, sev, sev_c, strat) in enumerate(rows, start=1):
        for j, val in enumerate((risk, sev, strat)):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if i % 2 else WHITE
            cell.margin_left = cell.margin_right = In(0.07)
            cell.margin_top = cell.margin_bottom = In(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.02
            r = p.add_run(); r.text = val
            r.font.name = SANS
            if j == 0:
                r.font.size = Pt(10.6); r.font.bold = True; r.font.color.rgb = NAVY
            elif j == 1:
                r.font.size = Pt(10.2); r.font.bold = True; r.font.color.rgb = sev_c
                p.alignment = PP_ALIGN.CENTER
            else:
                r.font.size = Pt(10.0); r.font.bold = False; r.font.color.rgb = INK

    rect(sl, tx, 6.24, tw_, 0.58, fill=PANEL2, line=NAVY, lw=1.0)
    tb(sl, tx, 6.24, tw_, 0.58, [{"runs": [
        ("Fallback ladder", 10.2, True, NAVY),
        ("     lose CSA → HMM map matching  ·  lose SVO → Tier B  ·  lose both → INS + NHC + "
         "ZUPT, already inside the 10 % gate. No single failure leaves us with nothing to show.",
         9.6, False, INK)], "line": 1.0}], anchor="m", pad=0.12)


# =============================================================== build
def main():
    prs = Presentation(SRC)
    s = prs.slides
    slide1(s[0]); slide2(s[1]); slide3(s[2]); slide4(s[3])
    v2.slide5(s[4]); v2.slide6(s[5])

    lst = prs.slides._sldIdLst
    ids = list(lst)
    prs.part.drop_rel(ids[6].get(qn("r:id")))
    lst.remove(ids[6])

    prs.save(DST)
    print("wrote", DST, "-", len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    main()
