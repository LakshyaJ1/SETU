"""Fill the SIH 2026 Idea template with the SETU submission.

Preserves every template element exactly: the SIH 2026 logo, the team-name oval,
the blue footer rail, footer text, slide numbers, the titles' Times New Roman voice
and the six-slide order. Only content is replaced. The mandated "idea details
pointers" are kept verbatim and promoted to be each section's own heading, so
compliance and hierarchy hold at the same time.

    python assets/build_deck.py
"""
import os
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "SIH2026-IDEA-Presentation-Format.pptx")
DST = os.path.join(ROOT, "SETU-SIH2026-Idea-Presentation.pptx")

# ---------------------------------------------------------------- design tokens
NAVY = RGBColor(0x14, 0x37, 0x5F)
BLUE = RGBColor(0x0C, 0x6D, 0xB5)
SAFFRON = RGBColor(0xE8, 0x76, 0x1F)
GREEN = RGBColor(0x1B, 0x8A, 0x4C)
RED = RGBColor(0xB3, 0x26, 0x1E)
INK = RGBColor(0x14, 0x21, 0x2B)
MUTED = RGBColor(0x5B, 0x6E, 0x78)
FIELD = RGBColor(0x9A, 0xA9, 0xB2)
RULE = RGBColor(0xCD, 0xD8, 0xDE)
PANEL = RGBColor(0xF3, 0xF7, 0xF9)
PANEL2 = RGBColor(0xE8, 0xEF, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xD6, 0xE3, 0xEC)

SANS, SERIF, MONO = "Arial", "Times New Roman", "Consolas"

TOP, BOT = 1.32, 6.82
L, R = 0.42, 12.91
W = R - L
GUT = 0.22
C3 = (W - 2 * GUT) / 3
X1, X2, X3 = L, L + C3 + GUT, L + 2 * (C3 + GUT)


# ---------------------------------------------------------------- primitives
def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = sl.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.text_frame.word_wrap = True
    return s


def hrule(sl, x, y, w, color=RULE, lw=0.9):
    ln = sl.shapes.add_connector(1, In(x), In(y), In(x + w), In(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(lw)
    return ln


def arrow(sl, x, y, w, h, color=RULE, shape=MSO_SHAPE.RIGHT_ARROW):
    s = sl.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    s.shadow.inherit = False
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def tb(sl, x, y, w, h, paras, anchor="t", pad=0.05):
    box = sl.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = In(pad)
    tf.margin_top = tf.margin_bottom = In(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if p.get("align"):
            para.alignment = p["align"]
        if p.get("sb") is not None:
            para.space_before = Pt(p["sb"])
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("line"):
            para.line_spacing = p["line"]
        for spec in p["runs"]:
            text, size, bold, color = spec[0], spec[1], spec[2], spec[3]
            r = para.add_run(); r.text = text
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = spec[4] if len(spec) > 4 else SANS
    return box


def picture(sl, x, y, w, img):
    iw, ih = Image.open(img).size
    h = w * ih / iw
    sl.shapes.add_picture(img, In(x), In(y), In(w), In(h))
    return y + h


def heading(sl, x, y, w, text, size=13.5):
    """A mandated template pointer, set as the section's own heading."""
    tb(sl, x, y, w, 0.32, [{"runs": [(text, size, True, NAVY)], "line": 0.94}], pad=0.0)
    hrule(sl, x, y + 0.29, w, NAVY, 1.4)
    return y + 0.42


def row(tag, body, tag_col=NAVY, size=9.5, tag_size=None, sa=4):
    runs = []
    if tag:
        runs.append((tag + "   ", tag_size or size, True, tag_col))
    runs.append((body, size, False, INK))
    return {"runs": runs, "sa": sa, "line": 1.02}


def kill(sl, *names):
    for sh in list(sl.shapes):
        if sh.name in names:
            sh._element.getparent().remove(sh._element)


def set_title(sl, text, size=None):
    for sh in sl.shapes:
        if sh.name.startswith("Title") and sh.has_text_frame:
            p = sh.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = text
                if size:
                    p.runs[0].font.size = Pt(size)
                for extra in p.runs[1:]:
                    extra.text = ""
            return


# =============================================================== SLIDE 1
def slide1(sl):
    kill(sl, "TextBox 9")
    for sh in sl.shapes:
        if sh.name.startswith("Subtitle") and sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = ""
            tf = sh.text_frame
            p = tf.paragraphs[1] if len(tf.paragraphs) > 1 else tf.paragraphs[0]
            r = p.add_run(); r.text = "SETU · INTELLIGENT DEAD RECKONING"
            r.font.size = Pt(23); r.font.bold = True
            r.font.color.rgb = NAVY; r.font.name = SERIF

    F = 19.0
    blank = ("____________", F, False, FIELD)
    rows = [
        [("Problem Statement ID –  ", F, True, NAVY), blank],
        [("Problem Statement Title-  ", F, True, NAVY),
         ("AI-ML based Intelligent Dead Reckoning system for seamless navigation",
          F, False, INK)],
        [("Theme-  ", F, True, NAVY), blank],
        [("PS Category-  ", F, True, NAVY), ("Software", F, False, INK)],
        [("Team ID-  ", F, True, NAVY), blank],
        [("Team Name (Registered on portal)-  ", F, True, NAVY), blank],
    ]
    tb(sl, L, 2.16, 6.62, 4.62,
       [{"runs": r, "sa": 9, "line": 1.0} for r in rows], pad=0.0)

    rect(sl, 0, 6.95, 13.333, 0.55, fill=BLUE)
    tb(sl, 0.45, 6.95, 12.43, 0.55, [{"runs": [
        ("SETU", 13, True, WHITE),
        ("   ·   Seamless Egomotion Tracking under Unavailable-GNSS", 11, False, PALE),
        ("      lane-level position in GNSS-denied environments, from a smartphone IMU alone",
         11, False, WHITE)], "align": PP_ALIGN.CENTER}], anchor="m")


# =============================================================== SLIDE 2
def slide2(sl):
    kill(sl, "TextBox 8")
    set_title(sl, "SETU — NAVIGATION WITHOUT GNSS", 32)

    tb(sl, L, TOP, W, 0.80, [{"runs": [
        ("Every phone-only system integrates acceleration twice, so its error grows as t². ",
         17, True, NAVY),
        ("SETU asks a different question — how far along this road am I — and answers it with "
         "three measurements whose error does not accumulate with time.", 17, False, INK)],
        "line": 1.06}], pad=0.0)
    hrule(sl, L, TOP + 0.84, W, RULE, 1.0)

    cy = 2.30
    HS = 11.5   # pointer headings sized to hold one line in a third-width column

    # ---- what it is
    y = heading(sl, X1, cy, C3, "Detailed explanation of the proposed solution", HS)
    tb(sl, X1, y, C3, 3.30, [
        {"runs": [("One C++ engine, two shells: Android app and edge daemon.",
                   9.2, False, MUTED)], "sa": 7, "line": 1.02},
        row("ACE", "solves phone→vehicle pitch, roll and yaw online. No user calibration; "
                   "re-aligns in under 2 s after a knock.", size=9.2),
        row("MSVR", "classifies motion and mount quality, so potholes, idling and phone "
                    "handling gate every other module.", size=9.2),
        row("SVO", "reads ground speed off the axle harmonics already present in the "
                   "accelerometer spectrum.", size=9.2),
        row("CTS", "reads absolute speed from lateral force ÷ yaw rate in every turn — "
                   "lean-corrected for two-wheelers.", size=9.2),
        row("CSA", "registers the gyro-derived path shape onto an offline OSM road manifold.",
            size=9.2),
        row("EFA", "magnetic, barometric and radio anchors, built by ordinary driving rather "
                   "than by a survey fleet.", size=9.2),
        row("IAF", "invariant EKF on SE₂(3) with learned noise, inside a road-graph particle "
                   "filter.", size=9.2),
        row("GQM", "per-satellite trust from raw GNSS — handover ramps, never switches.",
            size=9.2),
    ], pad=0.0)
    rect(sl, X1, 6.24, C3, 0.54, fill=PANEL, line=RULE)
    tb(sl, X1, 6.24, C3, 0.54, [{"runs": [
        ("3.6 MB", 10, True, NAVY, MONO), (" int8 models   ", 8.8, False, MUTED),
        ("2.7 ms", 10, True, NAVY, MONO), (" per tick   ", 8.8, False, MUTED),
        ("10 Hz", 10, True, NAVY, MONO), (" phone   ", 8.8, False, MUTED),
        ("200 Hz", 10, True, NAVY, MONO), (" edge   ·   fully offline", 8.8, False, MUTED)],
        "align": PP_ALIGN.CENTER, "line": 1.0}], anchor="m")

    # ---- why it works
    y = heading(sl, X2, cy, C3, "How it addresses the problem", HS)
    b = picture(sl, X2, y + 0.02, C3, os.path.join(HERE, "fig-error-law.png"))
    tb(sl, X2, b + 0.12, C3, 2.10, [
        row("Integrating", "a 0.1 m/s² bias gives ½·b·t² = 180 m in one minute. No network "
                           "repairs a double integrator.", tag_col=RED, size=9.4),
        row("Aiding velocity", "makes the error linear in distance — 31 m per km at the "
                               "3.1 %/km published for a smartphone in a real tunnel.",
            tag_col=SAFFRON, size=9.4),
        row("Registering position", "makes it bounded: error resets at every curvature "
                                    "landmark and field anchor, so it stops growing with time.",
            tag_col=GREEN, size=9.4),
    ], pad=0.0)

    # ---- what is new
    y = heading(sl, X3, cy, C3, "Innovation and uniqueness of the solution", HS)
    tb(sl, X3, y, C3, 2.28, [
        row("N1", "Speed from frequency, not integration. Chassis-vibration speed tracking "
                  "exists in the DSP literature and has never been fused into a navigation "
                  "filter.", size=9.0),
        row("N2", "Two-wheeler physics done right: v = g·sinφ / ωz. Omitting cosφ is a silent "
                  "4 % scale error — 40 m per km.", size=9.0),
        row("N3", "Scale-free map registration. Prior shape-matching work needs an odometer; "
                  "we solve for the scale itself.", size=9.0),
        row("N4", "IO-VNBD's synchronised CAN wheel speeds used as a training-only teacher "
                  "instead of noisy GPS speed.", size=9.0),
        row("N5", "Learned σ and learned Q inside one invariant filter, multi-hypothesis over "
                  "the road graph.", size=9.0),
    ], pad=0.0)
    b = picture(sl, X3, 4.82, C3, os.path.join(HERE, "fig-spectrogram.png"))
    tb(sl, X3, b + 0.06, C3, 0.34, [{"runs": [
        ("N1 evidence — axle orders rise continuously with speed while engine orders step down "
         "at each upshift. That step is a free training label.", 7.8, False, MUTED)],
        "line": 1.0}], pad=0.0)


# =============================================================== SLIDE 3
def slide3(sl):
    kill(sl, "TextBox 8")
    set_title(sl, "TECHNICAL APPROACH")

    heading(sl, L, TOP - 0.04, W, "Methodology and process for implementation")

    stages = [
        (L, 2.05, "SENSE", [
            ("Accel + gyro", "200–400 Hz"), ("Magnetometer", ""), ("Barometer", ""),
            ("GNSS raw", "incl. NavIC L5"), ("External IMU", "MEMS / FOG"),
        ], "no OBD-II · no vehicle bus"),
        (2.77, 1.95, "CONDITION", [
            ("STFT", "2.56 s / 0.1 s hop"), ("Harmonic ridge", "Viterbi"),
            ("Adaptive notch", "engine order"), ("Shock tagging", "potholes"),
            ("MSVR", "motion + mount"),
        ], "ACE solves the mount online"),
        (5.02, 3.45, "OBSERVE  ·  measurement generators", [
            ("SVO", "ground speed from axle harmonics"),
            ("CTS", "absolute speed from lateral force ÷ yaw rate"),
            ("CSA", "arc length from map shape, scale-free"),
            ("EFA", "magnetic / barometric / radio anchors"),
            ("NHC · ZUPT", "zero lateral, vertical, stopped velocity"),
            ("GQM", "per-satellite GNSS trust + Doppler"),
        ], "each emits ⟨ value, covariance, validity, timestamp ⟩"),
        (8.77, 2.30, "FUSE", [
            ("RI-EKF", "on SE₂(3) × θ, 23 states"),
            ("Learned σ", "per measurement"), ("Learned Q", "transformer"),
            ("RB-PF", "over the road graph"), ("64 / 256", "particles"),
        ], "χ² gating · stochastic cloning"),
        (11.37, 1.55, "DELIVER", [
            ("10 Hz", "phone"), ("200 Hz", "edge"), ("Lane-level", "pose"),
            ("Confidence tube", ""), ("gRPC · NMEA · ROS 2", ""),
        ], "no jump on re-acquire"),
    ]
    hy, hh, by, bh = 1.80, 0.30, 2.10, 1.74
    for x, w, name, rows, foot in stages:
        rect(sl, x, hy, w, hh, fill=NAVY)
        tb(sl, x, hy, w, hh, [{"runs": [(name, 10, True, WHITE)],
                               "align": PP_ALIGN.CENTER}], anchor="m", pad=0.03)
        rect(sl, x, by, w, bh, fill=PANEL, line=RULE)
        paras = []
        for a, b in rows:
            runs = [(a, 8.8, True, NAVY)]
            if b:
                runs.append(("   " + b, 8.4, False, MUTED))
            paras.append({"runs": runs, "sa": 3, "line": 1.0})
        tb(sl, x, by + 0.05, w, bh - 0.42, paras, pad=0.07)
        hrule(sl, x + 0.08, by + bh - 0.33, w - 0.16, RULE, 0.75)
        tb(sl, x, by + bh - 0.31, w, 0.29,
           [{"runs": [(foot, 7.8, True, BLUE)], "line": 1.0}], pad=0.07)

    for ax in (2.50, 4.75, 8.50, 11.10):
        arrow(sl, ax, by + bh / 2 - 0.11, 0.24, 0.22, RULE)

    # the two off-path inputs, as a matched pair of bands
    band_y = 4.06
    rect(sl, L, band_y, 4.45, 0.50, fill=PANEL2, line=RULE)
    tb(sl, L, band_y, 4.45, 0.50, [{"runs": [
        ("TRAINED OFF-DEVICE", 8.6, True, NAVY),
        ("   masked-IMU pretraining → privileged distillation from CAN → BPTT through the "
         "filter → int8 export", 8.2, False, INK)], "line": 1.0}], anchor="m", pad=0.07)

    rect(sl, 5.02, band_y, 6.05, 0.50, fill=PANEL2, line=RULE)
    tb(sl, 5.02, band_y, 6.05, 0.50, [{"runs": [
        ("OFFLINE MAP & ANCHOR BUNDLE", 8.6, True, NAVY),
        ("   OSM routing graph · ψ(s), κ(s) curvature LUT @ 1 m · DEM · magnetic and baro "
         "anchor DB · ≤250 MB per metro region", 8.2, False, INK)],
        "align": PP_ALIGN.CENTER, "line": 1.0}], anchor="m", pad=0.07)
    for ax in (6.55, 9.35):
        arrow(sl, ax, band_y - 0.17, 0.20, 0.16, RULE, MSO_SHAPE.UP_ARROW)

    heading(sl, L, 4.78, W, "Technologies to be used")
    tw = (W - 4 * 0.10) / 5
    groups = [
        ("Models & training", "PyTorch 2.5 · Lightning · Hydra · Polars · MLflow · "
                              "masked-IMU pretraining · privileged distillation from the CAN "
                              "bus · differentiable-filter BPTT"),
        ("On-device ML", "LiteRT + XNNPACK · ExecuTorch · int8 post-training quantisation, "
                         "QAT on the velocity head · export parity asserted in m/s, not in "
                         "tensor MSE"),
        ("Engine core", "C++20 · Eigen 3.4 · SE₂(3) Lie algebra · FlatBuffers + protobuf · "
                        "GTest with golden-replay determinism · one core, two shells"),
        ("Mobile", "Kotlin 2.0 · Jetpack Compose · MapLibre GL + PMTiles offline basemap · "
                   "SensorDirectChannel · GnssMeasurementsEvent · foreground service"),
        ("Edge & maps", "aarch64 static binary · gRPC / NMEA-0183 / ROS 2 · ADIS16505 and FOG "
                        "at 200 Hz · osmium · CartoDEM · tippecanoe → PMTiles"),
    ]
    for i, (name, body) in enumerate(groups):
        gx = L + i * (tw + 0.10)
        rect(sl, gx, 5.24, tw, 1.56, fill=PANEL, line=RULE)
        tb(sl, gx, 5.29, tw, 1.46, [
            {"runs": [(name, 9.6, True, NAVY)], "sa": 4, "line": 1.0},
            {"runs": [(body, 8.8, False, INK)], "line": 1.04}], pad=0.09)


# =============================================================== SLIDE 4
def slide4(sl):
    kill(sl, "TextBox 8")
    set_title(sl, "FEASIBILITY AND VIABILITY")

    fw = 3.92
    y = heading(sl, L, TOP - 0.04, fw, "Analysis of the feasibility of the idea")
    tb(sl, L, y, fw, 3.20, [
        row("Data already in hand.", "IO-VNBD gives 5,700 km over 98 h with the phone IMU and "
            "the CAN bus recorded simultaneously — both the input and a clean label.", size=9.4),
        row("The load-bearing physics is closed-form.", "v = a_lat / Ω and f = v / 2πR are "
            "algebra, not training. Only the residuals are learned, so the core carries no "
            "research risk.", size=9.4),
        row("Compute fits today's phones.", "3.6 MB of int8 models, 2.7 ms per 10 Hz tick, CPU "
            "only — no NPU or GPU delegate dependency across low-end SKUs.", size=9.4),
        row("The sensors are already there.", "Accelerometer, gyroscope, magnetometer, barometer "
            "and raw GNSS including NavIC L5 are exposed by standard Android APIs.", size=9.4),
        row("Maps are tractable offline.", "≤250 MB per metro region for the routing graph, the "
            "precomputed curvature lookup and the DEM. No connectivity in the runtime path.",
            size=9.4),
        row("Nothing to buy or install.", "No OBD-II dongle, no wheel encoder, no UWB, "
            "leaky-feeder or 5G roadside build-out — it runs on the phone already on the "
            "dashboard.", size=9.4),
    ], pad=0.0)

    # capability tiers: the honest statement of what runs on what
    ty = 4.94
    rect(sl, L, ty, fw, 1.38, fill=PANEL, line=RULE)
    tiers = [
        {"runs": [("Capability tiers — declared at runtime, never hidden", 9.6, True, NAVY)],
         "sa": 5, "line": 1.0},
        {"runs": [("A", 9.4, True, GREEN, MONO), ("   ≥200 Hz IMU", 8.8, True, INK),
                  ("   full spectral odometer   ", 8.6, False, MUTED),
                  ("0.5–1 %/km", 9.0, True, GREEN, MONO)], "sa": 3, "line": 1.0},
        {"runs": [("B", 9.4, True, BLUE, MONO), ("   50–200 Hz", 8.8, True, INK),
                  ("   axle orders 1–4 only   ", 8.6, False, MUTED),
                  ("1–2 %/km", 9.0, True, BLUE, MONO)], "sa": 3, "line": 1.0},
        {"runs": [("C", 9.4, True, SAFFRON, MONO), ("   10–50 Hz", 8.8, True, INK),
                  ("   no spectral odometer   ", 8.6, False, MUTED),
                  ("2–3 %/km", 9.0, True, SAFFRON, MONO)], "sa": 3, "line": 1.0},
        {"runs": [("IO-VNBD's smartphone stream is 10 Hz, so every result we report on it is "
                   "Tier C with the spectral odometer switched off.", 8.4, False, MUTED)],
         "sa": 0, "line": 1.02},
    ]
    tb(sl, L, ty + 0.05, fw, 1.28, tiers, pad=0.10)

    # risks, with the two mandated pointers used as the column headers
    tx, tw_ = 4.62, 8.29
    rows = [
        ("CSA is ambiguous on featureless or badly-mapped roads", "High", RED,
         "Landmark saliency scoring, top-k path hypotheses in the particle filter, honest σ from "
         "the Gauss–Newton Hessian. Fallback: classical HMM map matching, still inside the gate."),
        ("Axle harmonics too weak — smooth road, quiet EV, soft mount", "Medium", SAFFRON,
         "Mount-quality gate disables SVO when the mechanical path is decoupled; summation over "
         "8+ orders; learned σ down-weights it automatically. Fallback: Tier B, gate still met."),
        ("IO-VNBD's phone stream is 10 Hz, so it cannot validate SVO", "Known limit", BLUE,
         "Declared as capability tiers rather than hidden. IO-VNBD results are reported as "
         "Tier C with SVO off; SVO is validated on self-collected 400 Hz logs from week 2."),
        ("No two-wheeler data in IO-VNBD — four cars, no Indian roads", "Known limit", BLUE,
         "The lean relation is imposed as a physics loss rather than learned, so few samples are "
         "needed; a dedicated two-wheeler collection covers both mount types."),
        ("Phone sensor rate throttled below 200 Hz on some devices", "Medium", SAFFRON,
         "Tier detected at runtime from the achieved rate; SensorDirectChannel where available; "
         "verified across high, mid and low device classes."),
        ("Uneven OpenStreetMap geometry on Indian roads", "Medium", SAFFRON,
         "Cross-track covariance scaled by an OSM confidence heuristic, an explicit off-road "
         "particle so a missing road cannot force a wrong match, and a provider adapter."),
    ]
    shp = sl.shapes.add_table(len(rows) + 1, 3, In(tx), In(TOP - 0.04), In(tw_), In(4.9))
    t = shp.table
    t.first_row = False; t.horz_banding = False
    for c, wd in zip(t.columns, (2.80, 0.92, 4.57)):
        c.width = In(wd)
    for j, htxt in enumerate(["Potential challenges and risks", "Severity",
                              "Strategies for overcoming these challenges"]):
        cell = t.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        cell.margin_left = cell.margin_right = In(0.07)
        cell.margin_top = cell.margin_bottom = In(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = htxt
        r.font.size = Pt(9.6); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = SANS
        if j == 1:
            p.alignment = PP_ALIGN.CENTER
    for i, (risk, sev, sev_c, strat) in enumerate(rows, start=1):
        for j, val in enumerate((risk, sev, strat)):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if i % 2 else WHITE
            cell.margin_left = cell.margin_right = In(0.07)
            cell.margin_top = cell.margin_bottom = In(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.0
            r = p.add_run(); r.text = val
            r.font.name = SANS
            if j == 0:
                r.font.size = Pt(8.8); r.font.bold = True; r.font.color.rgb = NAVY
            elif j == 1:
                r.font.size = Pt(8.6); r.font.bold = True; r.font.color.rgb = sev_c
                p.alignment = PP_ALIGN.CENTER
            else:
                r.font.size = Pt(8.6); r.font.bold = False; r.font.color.rgb = INK


# =============================================================== SLIDE 5
def slide5(sl):
    kill(sl, "TextBox 8")
    set_title(sl, "IMPACT AND BENEFITS")

    aw = 5.05
    y = heading(sl, L, TOP - 0.04, aw, "Potential impact on the target audience")
    tb(sl, L, y, aw, 2.90, [
        row("Quick-commerce and delivery riders", "— the marker keeps moving through underpasses "
            "and basement parking, so the next turn is announced before the exit, not after it.",
            size=9.6),
        row("Logistics and fleet operators", "— mileage through tunnels and yards stops being "
            "guesswork; ETA and proof of delivery survive GNSS gaps, with no dongle to fit or "
            "maintain.", size=9.6),
        row("Ride-hailing drivers and riders", "— no frozen or teleporting car in an urban canyon, "
            "and pickup points that resolve to the right side of the road.", size=9.6),
        row("Emergency responders", "— ambulance and fire crews keep a live position inside "
            "tunnels, metro underpasses and multi-level structures, exactly where dispatch loses "
            "them today.", size=9.6),
        row("Two-wheeler riders", "— the largest vehicle class on Indian roads gets a lean-aware "
            "solution instead of one built for cars, and less time spent looking at a stalled "
            "screen.", size=9.6),
        row("Defence and critical logistics", "— a wholly self-contained position that degrades "
            "gracefully under jamming or spoofing instead of failing.", size=9.6),
    ], pad=0.0)

    sy = 4.52
    rect(sl, L, sy, aw, 2.06, fill=PANEL, line=RULE)
    tb(sl, L, sy + 0.05, aw, 1.96, [
        {"runs": [("Where it matters most", 10.4, True, NAVY)], "sa": 6, "line": 1.0},
        row("Long tunnels and metro underpasses", "— GNSS is absent, and the road is "
            "topologically one-dimensional, so the map does most of the work.", size=9.2),
        row("Multi-level parking", "— the helical ramp is a uniquely-shaped signature and the "
            "barometer resolves the floor to within one level.", size=9.2),
        row("Dense urban canyons", "— position is unusable but Doppler velocity survives on "
            "three or four satellites, where others discard everything.", size=9.2),
        row("Forested and valley highways", "— partial, intermittent GNSS, which the trust ramp "
            "absorbs without ever switching mode.", size=9.2, sa=0),
    ], pad=0.10)

    bx = 5.82
    bwid = W - (bx - L)
    bw = (bwid - GUT) / 2
    y2 = heading(sl, bx, TOP - 0.04, bwid,
                 "Benefits of the solution (social, economic, environmental, etc.)")
    blocks = [
        ("Social", GREEN,
         "Safer riding: a correct instruction on time means fewer glances at the screen at the "
         "moment a rider can least afford one. Emergency crews stay locatable underground. "
         "Accurate navigation stops being a feature of expensive cars and becomes a property of "
         "any phone on any dashboard."),
        ("Economic", BLUE,
         "Zero marginal hardware — it runs on phones riders already own. No OBD dongle per "
         "vehicle, no factory-fitted INS, no UWB or leaky-feeder installation per tunnel. Fewer "
         "missed exits, re-routes and failed delivery attempts per shift, on the same fleet."),
        ("Environmental", SAFFRON,
         "A missed exit is wasted kilometres, wasted fuel and avoidable emissions; removing the "
         "cause removes all three. And because the solution is entirely on-device, no roadside "
         "infrastructure has to be manufactured, powered or maintained."),
        ("Strategic", NAVY,
         "Resilient to jamming and spoofing by construction, since every SETU channel is "
         "self-contained. Uses NavIC L5 where Indian devices expose it, runs fully offline, and "
         "keeps trajectory data on the phone rather than in a cloud."),
    ]
    for i, (name, col, body) in enumerate(blocks):
        gx = bx + (i % 2) * (bw + GUT)
        gy = y2 + (i // 2) * 1.78
        rect(sl, gx, gy, bw, 1.62, fill=PANEL, line=RULE)
        tb(sl, gx, gy + 0.05, bw, 1.52, [
            {"runs": [(name, 11, True, col)], "sa": 4, "line": 1.0},
            {"runs": [(body, 9.0, False, INK)], "line": 1.04}], pad=0.10)

    # the benefit, quantified
    qy = y2 + 3.62
    cells = [
        ("Pure inertial, 60 s", "180 m", RED, "½·b·t² with a 0.1 m/s² bias — the marker is a "
                                              "block away"),
        ("Best published smartphone", "31 m", SAFFRON, "3.1 % of distance per km, INS + NHC "
                                                       "in a real tunnel"),
        ("SETU design target", "8–15 m", GREEN, "0.8–1.5 % per km, anchored to the map and to "
                                                "the field"),
    ]
    cw = (bwid - 2 * 0.16) / 3
    for i, (label, num, col, note) in enumerate(cells):
        gx = bx + i * (cw + 0.16)
        rect(sl, gx, qy, cw, 1.08, fill=WHITE, line=col, lw=1.1)
        tb(sl, gx, qy + 0.04, cw, 1.00, [
            {"runs": [(label, 8.6, True, MUTED)], "sa": 1, "line": 1.0},
            {"runs": [(num, 21, True, col)], "sa": 2, "line": 0.92},
            {"runs": [(note, 8.2, False, INK)], "line": 1.02}], pad=0.09)
    tb(sl, bx, qy + 1.14, bwid, 0.26, [{"runs": [
        ("Horizontal drift over a 1 km GNSS blackout at 60 km/h. The SIH ceiling is 10 % of "
         "distance travelled — 100 m.", 8.2, False, MUTED)], "line": 1.0}], pad=0.0)


# =============================================================== SLIDE 6
def slide6(sl):
    kill(sl, "TextBox 8")
    set_title(sl, "RESEARCH AND REFERENCES")

    heading(sl, L, TOP - 0.04, W, "Details / Links of the reference and research work")

    G = [
        ("Dataset and benchmark", [
            ("IO-VNBD: Inertial & Odometry Benchmark Dataset",
             "Onyekpe, Palade, Kanarachos, Szkolnik · Data in Brief 35:106885, 2021 · "
             "github.com/onyekpeu/IO-VNBD · arXiv:2005.01701"),
            ("WhONet: Wheel Odometry Neural Network",
             "Eng. Applications of AI, 2021 · arXiv:2104.02581 — our privileged upper bound"),
            ("R-WhONet: recalibration via transfer learning",
             "arXiv:2209.05877 — evidence for cross-vehicle domain shift"),
        ]),
        ("Learned inertial odometry", [
            ("Deep Learning for Inertial Positioning: A Survey",
             "Chen & Pan · IEEE T-ITS, 2024 · arXiv:2303.03757"),
            ("Inertial Navigation Meets Deep Learning — a survey",
             "Cohen & Klein · 2024 · arXiv:2307.00014"),
            ("AI-IMU Dead-Reckoning — invariant EKF, learned noise",
             "Brossard, Barrau & Bonnabel · IEEE T-IV, 2020 · arXiv:1904.06064"),
            ("TLIO: Tight Learned Inertial Odometry",
             "Liu et al. · IEEE RA-L, 2020 · arXiv:2007.01867 — learned covariance into a filter"),
            ("OdoNet: untethered speed aiding, no odometer",
             "arXiv:2109.03091 · 68 % error reduction over NHC alone"),
            ("CarSpeedNet: accelerometer-only speed",
             "arXiv:2401.07468 · <0.72 m/s ≈ 4.3 % — the accuracy to beat"),
            ("PiDR: Physics-Informed Inertial Dead Reckoning",
             "Sahoo & Klein · 2026 · arXiv:2601.03040"),
        ]),
        ("Learned and adaptive fusion", [
            ("KalmanNet: neural-network-aided Kalman filtering",
             "Revach, Shlezinger, van Sloun & Eldar · IEEE TSP, 2022"),
            ("A-KIT: Adaptive Kalman-Informed Transformer",
             "Cohen & Klein · arXiv:2401.09987 · github.com/ansfl/A-KIT — learned process noise"),
            ("Differentiable particle filters, semi-supervised",
             "arXiv:2011.05748"),
        ]),
        ("Map matching and kinematic constraints", [
            ("Hidden Markov Map Matching Through Noise and Sparseness",
             "Newson & Krumm · ACM SIGSPATIAL, 2009 — the map-matching baseline"),
            ("Heading–length sequence matching for localisation",
             "arXiv:2005.13704 — nearest prior art to CSA; assumes a known speed source"),
            ("Map-aided dead reckoning from OBD speed",
             "arXiv:1611.07910 — map-aided DR with a known odometer"),
            ("Lever-arm accuracy of the non-holonomic constraint",
             "Zhang & Hu · IEEE TVT, 2020"),
            ("NHC-assisted GNSS/SINS with motion-state CNN",
             "GPS Solutions, 2023 · doi 10.1007/s10291-023-01483-9"),
        ]),
        ("Spectral speed and field anchors", [
            ("Vehicle Speed Tracking Using Chassis Vibrations",
             "Linköping University, 2016 — axle-order tracking; motivates tunnel navigation"),
            ("Accelerometer-based wheel odometer",
             "Sensors, 2021 · PMC7918720"),
            ("Magnetic field vehicle positioning in a long tunnel",
             "Applied Sciences 11(24):11641, 2021 · doi 10.3390/app112411641"),
            ("Vehicle positioning in tunnel environments: a review",
             "Complex & Intelligent Systems, 2025 · doi 10.1007/s40747-024-01744-1"),
            ("Barometric floor detection on smartphones",
             "Sensors, 2015 · PMC4431287 — basis for the parking-level anchor"),
        ]),
        ("GNSS integrity and platform", [
            ("Vehicle positioning in underground space using a smartphone",
             "ISPRS Archives XLVI-3/W1, 2022 — the 3.1 %/km INS+NHC tunnel reference"),
            ("GNSS multipath detection using machine learning",
             "Hsu · IEEE ITSC, 2017"),
            ("Android GnssMeasurement HAL — raw GNSS observables",
             "source.android.com/reference/hal/struct_gnss_measurement"),
            ("u-blox Untethered Dead Reckoning (UDR)",
             "u-blox.com/en/technologies/udr-untethered-dead-reckoning — commercial reference"),
        ]),
    ]
    for x, groups in zip((X1, X2, X3), ([G[0], G[1]], [G[2], G[3]], [G[4], G[5]])):
        paras = []
        for gi, (gname, refs) in enumerate(groups):
            paras.append({"runs": [(gname.upper(), 9.2, True, BLUE)],
                          "sb": 0 if gi == 0 else 11, "sa": 4, "line": 1.0})
            for title, src in refs:
                paras.append({"runs": [(title, 8.8, True, INK)], "sa": 0, "line": 1.0})
                paras.append({"runs": [(src, 8.0, False, MUTED)], "sa": 5, "line": 1.0})
        tb(sl, x, 1.78, C3, 4.78, paras, pad=0.0)

    hrule(sl, L, 6.60, W, RULE, 1.0)
    tb(sl, L, 6.63, W, 0.26, [{"runs": [
        ("A 66-entry annotated bibliography, with a prior-art positioning table for each of the "
         "five mechanisms, is maintained alongside the implementation documentation.",
         8.2, False, MUTED)], "line": 1.0}], pad=0.0)


# =============================================================== build
def main():
    prs = Presentation(SRC)
    sl = prs.slides
    slide1(sl[0]); slide2(sl[1]); slide3(sl[2])
    slide4(sl[3]); slide5(sl[4]); slide6(sl[5])

    # the template's own instructions cap the deck at six slides including the
    # title, and authorise deleting the "Important Instructions" page before upload.
    lst = prs.slides._sldIdLst
    ids = list(lst)
    prs.part.drop_rel(ids[6].get(qn("r:id")))
    lst.remove(ids[6])

    prs.save(DST)
    print("wrote", DST, "-", len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    main()
